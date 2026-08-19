"""Convert screenshots, photos, and PDFs into editable Word docs and images."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
import pdfplumber
import pypdfium2 as pdfium
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageFilter, ImageOps
from rapidocr_onnxruntime import RapidOCR


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
PDF_EXTS = {".pdf"}
MAX_PAGE_IN = 21.5
MIN_PAGE_IN = 4.0
RENDER_SCALE = 2.0  # ~144–150 dpi for typical PDFs


@dataclass
class Box:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    size: float = 12.0
    font: str = ""
    bold: bool = False

    @property
    def cx(self) -> float:
        return (self.x0 + self.x1) / 2

    @property
    def cy(self) -> float:
        return (self.y0 + self.y1) / 2

    @property
    def w(self) -> float:
        return max(self.x1 - self.x0, 1.0)

    @property
    def h(self) -> float:
        return max(self.y1 - self.y0, 1.0)


@dataclass
class Line:
    boxes: list[Box]
    align: str = "left"

    @property
    def text(self) -> str:
        return join_boxes(self.boxes)

    @property
    def x0(self) -> float:
        return min(b.x0 for b in self.boxes)

    @property
    def y0(self) -> float:
        return min(b.y0 for b in self.boxes)

    @property
    def x1(self) -> float:
        return max(b.x1 for b in self.boxes)

    @property
    def y1(self) -> float:
        return max(b.y1 for b in self.boxes)

    @property
    def size(self) -> float:
        return max(b.size for b in self.boxes)

    @property
    def bold(self) -> bool:
        return any(b.bold for b in self.boxes) or is_heading_text(self.text)


@dataclass
class TableBlock:
    rows: list[list[str]]
    x0: float
    y0: float
    x1: float
    y1: float


@dataclass
class PageContent:
    width: float
    height: float
    lines: list[Line] = field(default_factory=list)
    tables: list[TableBlock] = field(default_factory=list)
    image_path: Path | None = None
    source: str = "ocr"
    raw_text: str = ""


class ConversionError(Exception):
    pass


_OCR: RapidOCR | None = None


def get_ocr() -> RapidOCR:
    global _OCR
    if _OCR is None:
        _OCR = RapidOCR()
    return _OCR


def warmup() -> None:
    get_ocr()


def has_devanagari(text: str) -> bool:
    return any("\u0900" <= ch <= "\u097F" for ch in text)


def is_heading_text(text: str) -> bool:
    t = text.strip()
    if not t or len(t) > 80:
        return False
    letters = [c for c in t if c.isalpha()]
    if letters and sum(c.isupper() for c in letters) / len(letters) > 0.82 and len(letters) > 3:
        return True
    return False


def join_boxes(boxes: list[Box]) -> str:
    if not boxes:
        return ""
    parts: list[str] = []
    for i, box in enumerate(boxes):
        t = box.text.strip()
        if not t:
            continue
        if not parts:
            parts.append(t)
            continue
        prev = parts[-1]
        if t[0] in ",.;:!?%)]}'\"”’":
            parts[-1] = prev + t
        elif prev[-1] in "([{\"'“‘/-":
            parts[-1] = prev + t
        else:
            parts.append(t)
    return " ".join(parts)


def group_lines(boxes: list[Box], y_tol: float = 0.55) -> list[Line]:
    if not boxes:
        return []
    items = sorted(boxes, key=lambda b: (b.cy, b.x0))
    clusters: list[list[Box]] = []
    for box in items:
        placed = False
        if clusters:
            last = clusters[-1]
            lcy = sum(b.cy for b in last) / len(last)
            lh = sum(b.h for b in last) / len(last)
            if abs(box.cy - lcy) <= max(box.h, lh) * y_tol:
                last.append(box)
                placed = True
        if not placed:
            clusters.append([box])
    lines: list[Line] = []
    for cluster in clusters:
        cluster.sort(key=lambda b: b.x0)
        lines.append(Line(boxes=cluster))
    return lines


def guess_align(line: Line, page_w: float) -> str:
    left_gap = line.x0
    right_gap = page_w - line.x1
    width = line.x1 - line.x0
    if width < page_w * 0.72:
        if abs(left_gap - right_gap) < page_w * 0.12 and left_gap > page_w * 0.18:
            return "center"
        if right_gap < page_w * 0.08 and left_gap > page_w * 0.22:
            return "right"
    return "left"


def enhance_for_ocr(im: Image.Image) -> Image.Image:
    rgb = im.convert("RGB")
    w, h = rgb.size
    long_side = max(w, h)
    if long_side < 1100:
        scale = 1400 / long_side
        rgb = rgb.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
    gray = ImageOps.autocontrast(rgb.convert("L"), cutoff=1)
    gray = gray.filter(ImageFilter.UnsharpMask(radius=1.2, percent=140, threshold=2))
    return gray.convert("RGB")


def ocr_image(im: Image.Image) -> list[Box]:
    engine = get_ocr()
    work = enhance_for_ocr(im)
    sx = im.size[0] / work.size[0]
    sy = im.size[1] / work.size[1]
    buf = io.BytesIO()
    work.save(buf, format="PNG")
    result, _ = engine(buf.getvalue())
    boxes: list[Box] = []
    if not result:
        return boxes
    for item in result:
        if not item or len(item) < 2:
            continue
        poly, text = item[0], item[1]
        score = item[2] if len(item) > 2 else 1.0
        if not text or not str(text).strip():
            continue
        try:
            if float(score) < 0.35:
                continue
        except (TypeError, ValueError):
            pass
        xs = [p[0] * sx for p in poly]
        ys = [p[1] * sy for p in poly]
        x0, x1 = min(xs), max(xs)
        y0, y1 = min(ys), max(ys)
        h = max(y1 - y0, 8)
        boxes.append(
            Box(
                text=str(text).strip(),
                x0=x0,
                y0=y0,
                x1=x1,
                y1=y1,
                size=max(8.0, h * 0.72),
                bold=is_heading_text(str(text)),
            )
        )
    return boxes


def render_pdf_page(pdf_path: Path, index: int, scale: float = RENDER_SCALE) -> Image.Image:
    doc = pdfium.PdfDocument(str(pdf_path))
    try:
        page = doc[index]
        bitmap = page.render(scale=scale)
        return bitmap.to_pil().convert("RGB")
    finally:
        doc.close()


def pdf_page_count(pdf_path: Path) -> int:
    doc = pdfium.PdfDocument(str(pdf_path))
    try:
        return len(doc)
    finally:
        doc.close()


def boxes_from_pdf_words(words: list[dict[str, Any]]) -> list[Box]:
    boxes: list[Box] = []
    for w in words:
        text = str(w.get("text") or "").strip()
        if not text:
            continue
        size = float(w.get("size") or 11)
        font = str(w.get("fontname") or "")
        bold = "bold" in font.lower() or "black" in font.lower() or is_heading_text(text)
        boxes.append(
            Box(
                text=text,
                x0=float(w["x0"]),
                y0=float(w["top"]),
                x1=float(w["x1"]),
                y1=float(w["bottom"]),
                size=size,
                font=font,
                bold=bold,
            )
        )
    return boxes


def extract_tables(page) -> list[TableBlock]:
    tables: list[TableBlock] = []
    try:
        found = page.find_tables() or []
    except Exception:
        found = []
    for tbl in found:
        try:
            data = tbl.extract()
        except Exception:
            continue
        if not data:
            continue
        rows = [[(cell or "").strip() for cell in row] for row in data]
        if not any(any(c for c in row) for row in rows):
            continue
        bbox = getattr(tbl, "bbox", None) or (0, 0, page.width, page.height)
        tables.append(
            TableBlock(
                rows=rows,
                x0=float(bbox[0]),
                y0=float(bbox[1]),
                x1=float(bbox[2]),
                y1=float(bbox[3]),
            )
        )
    return tables


def box_in_tables(box: Box, tables: list[TableBlock], pad: float = 2.0) -> bool:
    cx, cy = box.cx, box.cy
    for t in tables:
        if t.x0 - pad <= cx <= t.x1 + pad and t.y0 - pad <= cy <= t.y1 + pad:
            return True
    return False


def page_from_image(im: Image.Image, image_path: Path | None = None) -> PageContent:
    boxes = ocr_image(im)
    lines = group_lines(boxes)
    w, h = im.size
    for line in lines:
        line.align = guess_align(line, w)
    raw = "\n".join(ln.text for ln in lines if ln.text)
    return PageContent(
        width=float(w),
        height=float(h),
        lines=lines,
        image_path=image_path,
        source="ocr",
        raw_text=raw,
    )


def page_from_pdf(pdf_path: Path, index: int, image_path: Path) -> PageContent:
    im = render_pdf_page(pdf_path, index)
    im.save(image_path, format="PNG")

    digital_ok = False
    tables: list[TableBlock] = []
    boxes: list[Box] = []
    page_w = float(im.size[0])
    page_h = float(im.size[1])

    try:
        with pdfplumber.open(str(pdf_path)) as pdf:
            page = pdf.pages[index]
            page_w_pt = float(page.width)
            page_h_pt = float(page.height)
            sx = im.size[0] / page_w_pt
            sy = im.size[1] / page_h_pt
            words = page.extract_words(
                x_tolerance=2,
                y_tolerance=3,
                keep_blank_chars=False,
                extra_attrs=["fontname", "size"],
            )
            raw_chars = "".join((w.get("text") or "") for w in words)
            tables_pt = extract_tables(page)
            for t in tables_pt:
                tables.append(
                    TableBlock(
                        rows=t.rows,
                        x0=t.x0 * sx,
                        y0=t.y0 * sy,
                        x1=t.x1 * sx,
                        y1=t.y1 * sy,
                    )
                )
            if len(raw_chars.strip()) >= 25:
                digital_ok = True
                for b in boxes_from_pdf_words(words):
                    boxes.append(
                        Box(
                            text=b.text,
                            x0=b.x0 * sx,
                            y0=b.y0 * sy,
                            x1=b.x1 * sx,
                            y1=b.y1 * sy,
                            size=max(8.0, b.size * ((sx + sy) / 2) * 0.75),
                            font=b.font,
                            bold=b.bold,
                        )
                    )
    except Exception:
        digital_ok = False

    source = "pdf-text"
    if not digital_ok:
        boxes = ocr_image(im)
        source = "ocr"

    if tables:
        boxes = [b for b in boxes if not box_in_tables(b, tables)]

    lines = group_lines(boxes)
    for line in lines:
        line.align = guess_align(line, float(im.size[0]))
    raw = "\n".join(ln.text for ln in lines if ln.text)
    if tables:
        extra = []
        for t in tables:
            extra.append(
                "\n".join(" | ".join(c for c in row if c) for row in t.rows)
            )
        raw = (raw + "\n\n" + "\n\n".join(extra)).strip()

    return PageContent(
        width=page_w,
        height=page_h,
        lines=lines,
        tables=tables,
        image_path=image_path,
        source=source,
        raw_text=raw,
    )


def fit_page_size(px_w: float, px_h: float, dpi: float) -> tuple[float, float]:
    w_in = px_w / dpi
    h_in = px_h / dpi
    # Prefer A4 if the scan is roughly paper-shaped
    ratio = w_in / h_in
    a4 = 8.27 / 11.69
    letter = 8.5 / 11.0
    if 0.65 <= ratio <= 0.78 or abs(ratio - a4) < 0.06 or abs(ratio - letter) < 0.06:
        if h_in >= w_in:
            return 8.27, 11.69
        return 11.69, 8.27
    scale = 1.0
    if w_in > MAX_PAGE_IN or h_in > MAX_PAGE_IN:
        scale = min(MAX_PAGE_IN / w_in, MAX_PAGE_IN / h_in)
    if w_in < MIN_PAGE_IN and h_in < MIN_PAGE_IN:
        scale = max(MIN_PAGE_IN / w_in, MIN_PAGE_IN / h_in)
    return max(w_in * scale, 3.5), max(h_in * scale, 3.5)


def set_run_font(run, name: str, hindi: bool = False) -> None:
    run.font.name = "Nirmala UI" if hindi else name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Nirmala UI" if hindi else name)
    rFonts.set(qn("w:hAnsi"), "Nirmala UI" if hindi else name)
    rFonts.set(qn("w:cs"), "Nirmala UI")
    rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def prevent_table_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def add_table(doc: Document, table: TableBlock) -> None:
    rows = table.rows
    cols = max(len(r) for r in rows)
    if cols == 0:
        return
    tbl = doc.add_table(rows=len(rows), cols=cols)
    tbl.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            text = row[j] if j < len(row) else ""
            cell = tbl.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            hindi = has_devanagari(text)
            set_run_font(run, "Calibri", hindi)
            run.font.size = Pt(10.5)
            if i == 0:
                run.bold = True
                set_cell_shading(cell, "1B4332")
                run.font.color.rgb = RGBColor(255, 248, 235)
            elif i % 2 == 1:
                set_cell_shading(cell, "F4EFE4")
        prevent_table_row_split(tbl.rows[i])
    doc.add_paragraph("")


def add_page_lines(doc: Document, page: PageContent, page_w_in: float, page_h_in: float) -> None:
    usable_w = page_w_in - 0.9
    sx = page_w_in / page.width
    sy = page_h_in / page.height
    median_size = 12.0
    sizes = [ln.size * ((sx + sy) / 2) * 0.85 for ln in page.lines if ln.text]
    if sizes:
        sizes.sort()
        median_size = sizes[len(sizes) // 2]

    blocks: list[tuple[float, str, Any]] = []
    for ln in page.lines:
        if ln.text:
            blocks.append((ln.y0, "line", ln))
    for tb in page.tables:
        blocks.append((tb.y0, "table", tb))
    blocks.sort(key=lambda x: x[0])

    prev_y1 = 0.0
    first = True
    for _, kind, item in blocks:
        if kind == "table":
            add_table(doc, item)
            prev_y1 = item.y1
            first = False
            continue
        line: Line = item
        gap_px = 0 if first else max(0.0, line.y0 - prev_y1)
        first = False
        prev_y1 = line.y1
        gap_pt = gap_px * sy * 72.0

        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(2)
        pf.space_before = Pt(min(max(gap_pt - 2, 0), 36))
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        left_in = max(0.0, line.x0 * sx - 0.35)
        if line.align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.left_indent = Inches(0)
        elif line.align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            pf.left_indent = Inches(0)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if left_in > 0.15:
                pf.left_indent = Inches(min(left_in, usable_w * 0.45))

        font_pt = max(8.0, min(line.size * ((sx + sy) / 2) * 0.82, 28.0))
        heading = line.bold or font_pt > median_size * 1.28
        run = p.add_run(line.text)
        hindi = has_devanagari(line.text)
        set_run_font(run, "Calibri", hindi)
        run.font.size = Pt(round(font_pt, 1))
        run.bold = heading
        if heading and font_pt >= median_size * 1.35:
            run.font.color.rgb = RGBColor(18, 35, 26)


def build_editable_docx(pages: list[PageContent], out_path: Path, title: str) -> None:
    doc = Document()
    core = doc.core_properties
    core.author = "Kagaz Studio"
    core.title = title

    for i, page in enumerate(pages):
        section = doc.sections[0] if i == 0 else doc.add_section()
        dpi = 144.0 if page.source.startswith("pdf") else (96.0 if max(page.width, page.height) < 1800 else 150.0)
        w_in, h_in = fit_page_size(page.width, page.height, dpi)
        section.page_width = Inches(w_in)
        section.page_height = Inches(h_in)
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.5)

        if not page.lines and not page.tables:
            note = doc.add_paragraph()
            run = note.add_run("Is page se readable text nahi mil paya. Visual copy mein original dikhega.")
            set_run_font(run, "Calibri")
            run.italic = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(90, 80, 60)
        else:
            add_page_lines(doc, page, w_in, h_in)

    doc.save(str(out_path))


def build_visual_docx(pages: list[PageContent], out_path: Path, title: str) -> None:
    doc = Document()
    doc.core_properties.author = "Kagaz Studio"
    doc.core_properties.title = f"{title} — visual copy"

    for i, page in enumerate(pages):
        section = doc.sections[0] if i == 0 else doc.add_section()
        if not page.image_path or not page.image_path.exists():
            continue
        with Image.open(page.image_path) as im:
            px_w, px_h = im.size
        dpi = 150.0
        w_in, h_in = fit_page_size(px_w, px_h, dpi)
        section.page_width = Inches(w_in)
        section.page_height = Inches(h_in)
        section.left_margin = Inches(0)
        section.right_margin = Inches(0)
        section.top_margin = Inches(0)
        section.bottom_margin = Inches(0)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        run.add_picture(str(page.image_path), width=Inches(w_in), height=Inches(h_in))

    doc.save(str(out_path))


def write_images_zip(pages: list[PageContent], out_path: Path) -> None:
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for i, page in enumerate(pages, start=1):
            if page.image_path and page.image_path.exists():
                zf.write(page.image_path, arcname=f"page-{i:02d}.png")


def convert_file(
    src: Path,
    work_dir: Path,
    progress_cb=None,
) -> dict[str, Any]:
    work_dir.mkdir(parents=True, exist_ok=True)
    pages_dir = work_dir / "pages"
    pages_dir.mkdir(exist_ok=True)
    ext = src.suffix.lower()
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", src.stem)[:60] or "document"

    def tick(pct: int, message: str) -> None:
        if progress_cb:
            progress_cb(pct, message)

    pages: list[PageContent] = []

    if ext in PDF_EXTS:
        try:
            n = pdf_page_count(src)
        except Exception as exc:
            raise ConversionError("PDF khul nahi payi. File corrupt ya password-locked ho sakti hai.") from exc
        if n == 0:
            raise ConversionError("PDF khali hai.")
        n = min(n, 40)
        for i in range(n):
            tick(8 + int(70 * i / n), f"Page {i + 1}/{n} padh rahe hain…")
            image_path = pages_dir / f"page-{i + 1:02d}.png"
            pages.append(page_from_pdf(src, i, image_path))
    elif ext in IMAGE_EXTS:
        tick(15, "Photo / screenshot padh rahe hain…")
        try:
            with Image.open(src) as raw:
                raw = ImageOps.exif_transpose(raw)
                im = raw.convert("RGB")
                if getattr(raw, "n_frames", 1) > 1 and ext in {".tif", ".tiff"}:
                    frames = []
                    for fi in range(min(raw.n_frames, 40)):
                        raw.seek(fi)
                        frames.append(ImageOps.exif_transpose(raw).convert("RGB"))
                else:
                    frames = [im]
        except Exception as exc:
            raise ConversionError("Image khul nahi payi. Dusri file try karein.") from exc
        for i, frame in enumerate(frames):
            tick(20 + int(60 * i / max(len(frames), 1)), f"OCR page {i + 1}/{len(frames)}…")
            image_path = pages_dir / f"page-{i + 1:02d}.png"
            frame.save(image_path, format="PNG")
            pages.append(page_from_image(frame, image_path))
    else:
        raise ConversionError("Sirf PDF, PNG, JPG, WEBP, BMP ya TIFF chalegi.")

    if not pages:
        raise ConversionError("Koi page process nahi ho paya.")

    tick(88, "Editable Word file bana rahe hain…")
    editable = work_dir / f"{stem}-editable.docx"
    visual = work_dir / f"{stem}-visual.docx"
    images_zip = work_dir / f"{stem}-pages.zip"
    build_editable_docx(pages, editable, stem)
    tick(93, "Visual Word copy bana rahe hain…")
    build_visual_docx(pages, visual, stem)
    write_images_zip(pages, images_zip)

    preview = []
    full_text_parts = []
    for i, page in enumerate(pages, start=1):
        preview.append(
            {
                "index": i,
                "text": page.raw_text,
                "source": page.source,
                "width": int(page.width),
                "height": int(page.height),
                "lines": len(page.lines),
                "tables": len(page.tables),
            }
        )
        if page.raw_text:
            full_text_parts.append(f"— Page {i} —\n{page.raw_text}")

    tick(100, "Ready")
    return {
        "pages": len(pages),
        "editable": editable.name,
        "visual": visual.name,
        "images_zip": images_zip.name,
        "preview": preview,
        "text": "\n\n".join(full_text_parts).strip(),
        "stem": stem,
    }
